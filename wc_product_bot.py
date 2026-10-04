#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
wc-product-bot
===============
Bulk-create WooCommerce products from an Excel sheet, using an existing
product as a template (for price, categories, images, product type, etc.)
and filling the title / slug / description / short description from
customizable text templates.

Usage:
    python wc_product_bot.py --config config.ini
    python wc_product_bot.py --config config.ini --dry-run
    python wc_product_bot.py --config config.ini --limit 1

See README.md for full setup instructions.
License: MIT (see LICENSE file).
"""

from __future__ import annotations

import argparse
import configparser
import re
import sys
import time
from pathlib import Path
from typing import Any

import openpyxl
import requests


# ---------------------------------------------------------------------------
# Config loading
# ---------------------------------------------------------------------------

class Config:
    def __init__(self, path: str):
        parser = configparser.ConfigParser()
        read_files = parser.read(path, encoding="utf-8")
        if not read_files:
            raise FileNotFoundError(f"Config file not found: {path}")

        wc = parser["woocommerce"]
        self.site_url: str = wc["site_url"].rstrip("/")
        self.consumer_key: str = wc["consumer_key"]
        self.consumer_secret: str = wc["consumer_secret"]
        self.api_base: str = f"{self.site_url}/wp-json/wc/v3"

        tpl = parser["template_product"]
        self.template_product_id: int = int(tpl["product_id"])
        self.publish_status: str = tpl.get("publish_status", "draft")
        self.copy_fields: list[str] = [
            f.strip() for f in tpl.get("copy_fields", "").split(",") if f.strip()
        ]

        xl = parser["excel"]
        self.excel_path: str = xl["path"]
        self.header_row: int = int(xl.get("header_row", "1"))
        self.status_column: str = xl.get("status_column", "").strip() or None
        self.done_value: str = xl.get("done_value", "done")
        self.error_value: str = xl.get("error_value", "error")

        if "columns" not in parser:
            raise ValueError(
                "Config is missing the [columns] section — this maps your "
                "Excel columns to placeholder names used in the templates."
            )
        # e.g. name = A  ->  self.columns["name"] = "A"
        self.columns: dict[str, str] = dict(parser["columns"])

        seo = parser["seo"] if "seo" in parser else {}
        self.set_focus_keyword: bool = str(seo.get("set_focus_keyword", "false")).lower() == "true"
        self.focus_keyword_meta_key: str = seo.get("focus_keyword_meta_key", "rank_math_focus_keyword")

        req = parser["request"] if "request" in parser else {}
        self.user_agent: str = req.get(
            "user_agent",
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        )
        self.delay_seconds: float = float(req.get("delay_seconds", "1.5"))
        self.timeout_seconds: float = float(req.get("timeout_seconds", "30"))

        tplf = parser["templates"] if "templates" in parser else {}
        self.title_template_path: str | None = tplf.get("title_file") or None
        self.description_template_path: str = tplf["description_file"]
        self.short_description_template_path: str | None = tplf.get("short_description_file") or None
        self.slug_template: str | None = tplf.get("slug_template") or None


def load_text_template(path: str) -> str:
    return Path(path).read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------

def load_rows(cfg: Config) -> tuple[Any, Any, list[dict]]:
    """
    Reads every data row from the Excel sheet and turns it into a dict keyed
    by the placeholder names from [columns] in the config, skipping rows
    already marked done in the status column (if one is configured).
    """
    wb = openpyxl.load_workbook(cfg.excel_path)
    ws = wb.active
    rows = []

    for row_idx in range(cfg.header_row + 1, ws.max_row + 1):
        data = {"row_idx": row_idx}
        any_value = False
        for placeholder, col_letter in cfg.columns.items():
            cell_val = ws[f"{col_letter}{row_idx}"].value
            if cell_val is not None and str(cell_val).strip() != "":
                any_value = True
            data[placeholder] = str(cell_val).strip() if cell_val is not None else ""

        if not any_value:
            continue  # every mapped column is empty -> treat as end of data / blank row

        if cfg.status_column:
            status_val = ws[f"{cfg.status_column}{row_idx}"].value
            if status_val == cfg.done_value:
                continue

        rows.append(data)

    return wb, ws, rows


def mark_status(cfg: Config, wb, ws, row_idx: int, text: str) -> None:
    if not cfg.status_column:
        return
    ws[f"{cfg.status_column}{row_idx}"] = text
    wb.save(cfg.excel_path)


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

def slugify(text: str) -> str:
    """Simple ASCII slugifier. Replace with your own logic if you need
    non-Latin slugs (WordPress does accept them)."""
    text = text.strip().lower()
    text = re.sub(r"\s+", "-", text)
    text = re.sub(r"[^a-z0-9\-]", "", text)
    text = re.sub(r"-+", "-", text).strip("-")
    return text


# ---------------------------------------------------------------------------
# WooCommerce REST API
# ---------------------------------------------------------------------------

class WooClient:
    def __init__(self, cfg: Config):
        self.cfg = cfg
        self.headers = {
            # Some hosts / security plugins block the default python-requests
            # User-Agent as "bot-like" traffic and return 403. Sending a
            # normal browser User-Agent avoids that on most shared hosting.
            "User-Agent": cfg.user_agent,
            "Accept": "application/json",
        }
        # Sending the key/secret as query-string params rather than an
        # Authorization header, because some hosts strip that header before
        # it reaches WordPress (also a common 403 cause). If your host is
        # fine with it, you can switch to `auth=(key, secret)` instead.
        self.auth_params = {
            "consumer_key": cfg.consumer_key,
            "consumer_secret": cfg.consumer_secret,
        }

    def get_template_product(self) -> dict:
        url = f"{self.cfg.api_base}/products/{self.cfg.template_product_id}"
        resp = requests.get(
            url, params=self.auth_params, headers=self.headers, timeout=self.cfg.timeout_seconds
        )
        resp.raise_for_status()
        return resp.json()

    def create_product(self, payload: dict) -> dict:
        url = f"{self.cfg.api_base}/products"
        resp = requests.post(
            url,
            json=payload,
            params=self.auth_params,
            headers=self.headers,
            timeout=self.cfg.timeout_seconds,
        )
        resp.raise_for_status()
        return resp.json()


# ---------------------------------------------------------------------------
# Payload building
# ---------------------------------------------------------------------------

def build_payload(cfg: Config, template: dict, data: dict,
                   title_tpl: str | None, desc_tpl: str, short_tpl: str | None) -> dict:
    new_title = title_tpl.format(**data).strip() if title_tpl else data.get("name", "")

    payload: dict[str, Any] = {
        "name": new_title,
        "description": desc_tpl.format(**data),
        "status": cfg.publish_status,
    }

    if short_tpl:
        payload["short_description"] = short_tpl.format(**data)

    if cfg.slug_template:
        payload["slug"] = slugify(cfg.slug_template.format(**data))

    # Copy whichever top-level fields the config says to carry over from
    # the template product (price, type, stock settings, etc.)
    for field in cfg.copy_fields:
        if field in template:
            payload[field] = template[field]

    if template.get("categories"):
        payload["categories"] = [{"id": c["id"]} for c in template["categories"]]

    if template.get("images"):
        payload["images"] = [{"id": img["id"]} for img in template["images"] if img.get("id")]

    if cfg.set_focus_keyword:
        payload.setdefault("meta_data", [])
        payload["meta_data"].append({"key": cfg.focus_keyword_meta_key, "value": new_title})

    return payload


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Bulk-create WooCommerce products from Excel.")
    parser.add_argument("--config", default="config.ini", help="Path to config.ini (default: config.ini)")
    parser.add_argument("--dry-run", action="store_true", help="Build payloads but don't call the API or touch the Excel file.")
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N pending rows (for testing).")
    args = parser.parse_args()

    cfg = Config(args.config)

    title_tpl = load_text_template(cfg.title_template_path) if cfg.title_template_path else None
    desc_tpl = load_text_template(cfg.description_template_path)
    short_tpl = load_text_template(cfg.short_description_template_path) if cfg.short_description_template_path else None

    wb, ws, rows = load_rows(cfg)
    if args.limit:
        rows = rows[: args.limit]

    if not rows:
        print("Nothing to do — every row is already marked done, or the sheet is empty.")
        return

    print(f"{len(rows)} product(s) to create.")

    client = WooClient(cfg)
    if args.dry_run:
        print("[dry-run] Skipping the real API call for the template product "
              "(copy_fields won't be populated in the printed payloads).")
        template: dict = {}
    else:
        template = client.get_template_product()
        print(f"Loaded template product: {template.get('name')!r}")

    for data in rows:
        try:
            payload = build_payload(cfg, template, data, title_tpl, desc_tpl, short_tpl)
            label = payload.get("name") or data.get("name") or f"row {data['row_idx']}"
            print(f"Creating: {label} ...")

            if args.dry_run:
                print("  [dry-run] payload built OK, not calling the API:")
                print(f"  name:  {payload.get('name')}")
                print(f"  slug:  {payload.get('slug')}")
                continue

            created = client.create_product(payload)
            print(f"  OK -> {created.get('name')}  ({created.get('permalink', '')})")
            mark_status(cfg, wb, ws, data["row_idx"], cfg.done_value)
        except Exception as exc:  # noqa: BLE001 - we want to log and keep going
            print(f"  FAILED ({exc})", file=sys.stderr)
            if not args.dry_run:
                mark_status(cfg, wb, ws, data["row_idx"], cfg.error_value)

        if not args.dry_run:
            time.sleep(cfg.delay_seconds)

    print("Done.")


if __name__ == "__main__":
    main()
